"""Flujo persistente XLSX de Ausentismos, sin inserciones durante el análisis."""

import json
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.core.permissions import AuthenticatedUser, authorize
from app.models import Ausentismo, ErrorImportacionAusentismo, LoteImportacionAusentismo
from app.schemas.ausentismos import AusentismoAnalisis
from app.services.audit import log_change
from app.services.ausentismos import analizar_importacion_ausentismos
from app.services.records import creation_metadata, get_active, mark_updated

MAX_XLSX_BYTES = 10 * 1024 * 1024
ESTADO_ANALIZADO = "ANALIZADO"
ESTADO_CONFIRMADO = "CONFIRMADO"


def leer_xlsx(contenido: bytes) -> tuple[list[object], list[list[object]]]:
    if not contenido:
        raise AppError("EMPTY_IMPORT_FILE", "El archivo XLSX está vacío.", 422)
    if len(contenido) > MAX_XLSX_BYTES:
        raise AppError("IMPORT_FILE_TOO_LARGE", "El archivo XLSX supera el tamaño máximo permitido.", 422)
    try:
        workbook = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise AppError("INVALID_XLSX", "El archivo no es un XLSX válido.", 422) from exc
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = list(next(iterator, ()))
        if not headers:
            raise AppError("EMPTY_IMPORT_FILE", "El archivo XLSX no contiene encabezados.", 422)
        rows = [list(row) for row in iterator if any(value is not None and str(value).strip() for value in row)]
        if not rows:
            raise AppError("EMPTY_IMPORT_FILE", "El archivo XLSX no contiene filas de datos.", 422)
        return headers, rows
    finally:
        workbook.close()


def _safe_row_data(row) -> str:
    cedula = row.cedula
    masked = None if cedula is None else ("*" * max(0, len(cedula) - 4) + cedula[-4:])
    return json.dumps({
        "cedula": masked, "fecha_inicio": row.fecha_inicio, "fecha_fin": row.fecha_fin,
        "tipo_ausentismo": row.tipo_ausentismo, "motivo": row.motivo,
    }, ensure_ascii=False)


def _save_issues(session: Session, lote: LoteImportacionAusentismo, analysis: AusentismoAnalisis, user: AuthenticatedUser) -> None:
    for row in analysis.filas:
        if row.estado == "VALIDA":
            continue
        code = "DUPLICADO" if row.estado == "DUPLICADA" else "VALIDACION"
        message = "Registro duplicado según persona, fechas y tipo de ausentismo." if code == "DUPLICADO" else " ".join(row.errores)
        session.add(ErrorImportacionAusentismo(
            id_error=str(uuid4()), lote_id=lote.id_lote, numero_fila=row.fila, codigo=code,
            mensaje=message, datos_fila=_safe_row_data(row), **creation_metadata(user.correo),
        ))


def _apply_counts(lote: LoteImportacionAusentismo, analysis: AusentismoAnalisis) -> None:
    lote.total_filas = analysis.total_filas
    lote.filas_validas = analysis.filas_validas
    lote.filas_con_error = analysis.filas_con_error
    lote.filas_duplicadas = analysis.filas_duplicadas


def analizar_archivo_ausentismos(
    session: Session, user: AuthenticatedUser, *, nombre_archivo: str, contenido: bytes, correlation_id: str,
) -> tuple[LoteImportacionAusentismo, AusentismoAnalisis]:
    """Persiste lote y sus incidencias, pero nunca crea Ausentismos."""
    authorize(user, "IMPORTACION", "create")
    authorize(user, "AUSENTISMO", "create")
    if not nombre_archivo.lower().endswith(".xlsx"):
        raise AppError("INVALID_XLSX", "El archivo debe tener extensión .xlsx.", 422)
    headers, rows = leer_xlsx(contenido)
    analysis = analizar_importacion_ausentismos(session, user, encabezados=headers, filas=rows)
    lote = LoteImportacionAusentismo(
        id_lote=str(uuid4()), nombre_archivo=nombre_archivo, contenido_archivo=contenido, usuario_id=user.id_usuario,
        estado=ESTADO_ANALIZADO, filas_importadas=0, **creation_metadata(user.correo),
    )
    _apply_counts(lote, analysis)
    session.add(lote)
    session.flush()
    _save_issues(session, lote, analysis, user)
    log_change(session, "lotes_importacion_ausentismo", lote.id_lote, "IMPORT", {}, _lote_snapshot(lote),
               user.correo, "Análisis de archivo XLSX", correlation_id)
    return lote, analysis


def _lote_snapshot(lote: LoteImportacionAusentismo) -> dict:
    return {
        "nombre_archivo": lote.nombre_archivo, "estado": lote.estado, "total_filas": lote.total_filas,
        "filas_validas": lote.filas_validas, "filas_con_error": lote.filas_con_error,
        "filas_duplicadas": lote.filas_duplicadas, "filas_importadas": lote.filas_importadas,
    }


def confirmar_lote_ausentismos(session: Session, user: AuthenticatedUser, lote_id: str, *, correlation_id: str) -> LoteImportacionAusentismo:
    """Reanaliza y persiste el lote como una unidad atómica."""
    authorize(user, "IMPORTACION", "create")
    authorize(user, "AUSENTISMO", "create")
    lote = get_active(session, LoteImportacionAusentismo, lote_id, LoteImportacionAusentismo.id_lote)
    if lote.estado == ESTADO_CONFIRMADO:
        raise AppError("IMPORT_ALREADY_CONFIRMED", "El lote ya fue confirmado.", 409)
    if lote.estado != ESTADO_ANALIZADO:
        raise AppError("INVALID_IMPORT_STATE", "El lote no está disponible para confirmación.", 409)
    headers, rows = leer_xlsx(lote.contenido_archivo)
    analysis = analizar_importacion_ausentismos(session, user, encabezados=headers, filas=rows)
    if not analysis.puede_confirmarse:
        raise AppError("IMPORT_NOT_CONFIRMABLE", "El lote tiene errores o duplicados y no puede confirmarse.", 422)
    before = _lote_snapshot(lote)
    try:
        with session.begin_nested():
            for row in analysis.filas:
                session.add(Ausentismo(
                    id_ausentismo=str(uuid4()), persona_id=row.persona_id, fecha_inicio=row.fecha_inicio,
                    fecha_fin=row.fecha_fin, tipo_ausentismo=row.tipo_ausentismo, motivo=row.motivo,
                    observacion=row.observacion, archivo_fuente=lote.nombre_archivo, registro_fuente=str(row.fila),
                    fecha_importacion=lote.fecha_creacion, usuario_importacion=user.correo,
                    **creation_metadata(user.correo),
                ))
            session.flush()
    except IntegrityError as exc:
        raise AppError("IMPORT_PERSISTENCE_FAILED", "No fue posible confirmar el lote; no se importó ninguna fila.", 409) from exc
    lote.estado = ESTADO_CONFIRMADO
    lote.filas_importadas = analysis.filas_validas
    mark_updated(lote, user.correo)
    log_change(session, "lotes_importacion_ausentismo", lote.id_lote, "IMPORT", before, _lote_snapshot(lote),
               user.correo, "Confirmación de importación XLSX", correlation_id)
    return lote


def listar_lotes_ausentismos(session: Session, user: AuthenticatedUser, *, limit: int, offset: int):
    authorize(user, "IMPORTACION", "read")
    authorize(user, "AUSENTISMO", "read")
    stmt = select(LoteImportacionAusentismo).where(LoteImportacionAusentismo.eliminado.is_(False))
    total = session.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    lots = session.scalars(stmt.order_by(LoteImportacionAusentismo.fecha_creacion.desc()).offset(offset).limit(limit)).all()
    return lots, total


def obtener_lote_ausentismos(session: Session, user: AuthenticatedUser, lote_id: str) -> LoteImportacionAusentismo:
    authorize(user, "IMPORTACION", "read")
    authorize(user, "AUSENTISMO", "read")
    return get_active(session, LoteImportacionAusentismo, lote_id, LoteImportacionAusentismo.id_lote)


def listar_errores_lote_ausentismos(session: Session, user: AuthenticatedUser, lote_id: str, *, limit: int, offset: int):
    obtener_lote_ausentismos(session, user, lote_id)
    stmt = select(ErrorImportacionAusentismo).where(
        ErrorImportacionAusentismo.lote_id == lote_id, ErrorImportacionAusentismo.eliminado.is_(False),
    )
    total = session.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    errors = session.scalars(stmt.order_by(ErrorImportacionAusentismo.numero_fila, ErrorImportacionAusentismo.id_error).offset(offset).limit(limit)).all()
    return errors, total
