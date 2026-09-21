"""Exportación XLSX privada de la bandeja de correos.

No conserva el XLSX de entrada: crea una vista actual de ``Correo`` y, cuando
corresponde, de ``SeguimientoCorreo``.  ``write_only`` evita cargar el libro
completo en memoria durante exportaciones extensas.
"""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from time import perf_counter
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.core.permissions import AuthenticatedUser, authorize
from app.core.time import ECUADOR_TZ
from app.models import Correo, SeguimientoCorreo
from app.services.audit import log_change
from app.services.correos import email_selection

MAX_CELL_CHARS = 32_767
# Se reserva un carácter para la protección reversible contra fórmulas.
CONTENT_CHUNK_CHARS = MAX_CELL_CHARS - 1
FOLLOW_UP_CHUNK_SIZE = 400
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_ESCAPE_MARKER = "\u200b"
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_TEXT_ALIGNMENT = Alignment(vertical="top", wrap_text=True)


@dataclass(frozen=True)
class EmailExport:
    path: str
    filename: str
    count: int
    duration_ms: int
    extended_content: bool


CORREOS_COLUMNS = (
    ("ID interno del correo", "id_correo"),
    ("MessageId", "id_externo_correo"),
    ("Fecha de recepción", "fecha_recibido"),
    ("Remitente", "remitente"),
    ("Destinatarios", "destinatarios"),
    ("CC", "cc"),
    ("Asunto", "asunto"),
    ("Importancia", "importancia"),
    ("Tiene adjuntos", "tiene_adjuntos"),
    ("Leído", "leido"),
    ("Categoría macro", "categoria_macro"),
    ("Categoría", "categoria_nombre"),
    ("Estado de categoría", "estado_categoria"),
    ("Regla disparadora", "regla_disparadora"),
    ("Estado de clasificación", "estado_clasificacion"),
    ("Estado del requerimiento", "estado_requerimiento"),
    ("Responsable de seguimiento", "responsable_seguimiento"),
    ("Origen", "archivo_fuente"),
    ("ID del lote", "lote_id"),
    ("Fecha de creación", "fecha_creacion"),
    ("Fecha de actualización", "fecha_actualizacion"),
)

FILTER_LABELS = {
    "estado": "Estado del requerimiento", "categoria": "Categoría macro", "texto": "Texto",
    "asunto": "Asunto", "remitente": "Remitente", "destinatario": "Destinatario",
    "message_id": "MessageId", "fecha_desde": "Fecha desde", "fecha_hasta": "Fecha hasta",
    "origen": "Origen", "importancia": "Importancia", "tiene_adjuntos": "Tiene adjuntos",
    "orden": "Orden",
}


def remove_export_file(path: str) -> None:
    try:
        os.unlink(path)
    except FileNotFoundError:
        pass


def _safe_text(value) -> str | None:
    if value is None:
        return None
    text = str(value)
    # Un marcador inicial protege valores que Excel interpretaría como fórmulas.
    # Si el valor original empieza por el propio marcador se duplica, de modo que
    # la transformación pueda revertirse sin ambigüedad.
    if text.startswith(_ESCAPE_MARKER) or text.startswith(_FORMULA_PREFIXES):
        return _ESCAPE_MARKER + text
    return text


def _cell(sheet, value, *, header: bool = False):
    cell = WriteOnlyCell(sheet, value=_safe_text(value))
    cell.alignment = _HEADER_ALIGNMENT if header else _TEXT_ALIGNMENT
    if header:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    else:
        cell.number_format = "@"
    return cell


def _header(sheet, headings: list[str]) -> None:
    sheet.append([_cell(sheet, heading, header=True) for heading in headings])
    sheet.freeze_panes = "A2"


def _configure_columns(sheet, widths: list[int]) -> None:
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _append_content(value, *, correo_id: str, field: str, extended_sheet) -> tuple[str | None, int]:
    if value is None:
        return None, 0
    text = str(value)
    if len(_safe_text(text) or "") <= MAX_CELL_CHARS:
        return text, 0
    total = (len(text) + CONTENT_CHUNK_CHARS - 1) // CONTENT_CHUNK_CHARS
    for part in range(total):
        start = part * CONTENT_CHUNK_CHARS
        extended_sheet.append([_cell(extended_sheet, correo_id), _cell(extended_sheet, field), _cell(extended_sheet, part + 1), _cell(extended_sheet, total), _cell(extended_sheet, text[start:start + CONTENT_CHUNK_CHARS])])
    return "[Contenido extendido: consulte Contenido_extenso]", total


def _active_filters(filters: dict) -> str:
    values = []
    for key, label in FILTER_LABELS.items():
        value = filters.get(key)
        if value not in (None, ""):
            values.append(f"{label}: {'Sí' if value is True else 'No' if value is False else value}")
    return "; ".join(values) if values else "Sin filtros"


def _follow_ups_for(session: Session, correo_ids: list[str]) -> dict[str, list[SeguimientoCorreo]]:
    if not correo_ids:
        return {}
    rows = session.scalars(
        select(SeguimientoCorreo).where(
            SeguimientoCorreo.correo_id.in_(correo_ids),
            SeguimientoCorreo.eliminado.is_(False),
        ).order_by(SeguimientoCorreo.correo_id, SeguimientoCorreo.fecha_seguimiento, SeguimientoCorreo.fecha_creacion, SeguimientoCorreo.id_seguimiento)
    )
    result: dict[str, list[SeguimientoCorreo]] = {}
    for row in rows:
        result.setdefault(row.correo_id, []).append(row)
    return result


def generate_email_export(
    session: Session,
    user: AuthenticatedUser,
    *,
    scope: str,
    filters: dict,
    include_body: bool,
    include_follow_ups: bool,
    correlation_id: str,
) -> EmailExport:
    """Genera y audita una exportación, con conteo y filas del mismo SELECT ordenado."""
    authorize(user, "CORREOS", "read")
    authorize(user, "CORREOS", "export")
    if include_body or include_follow_ups:
        authorize(user, "CORREOS", "sensitive")
    if scope not in {"filtered", "all"}:
        raise AppError("INVALID_EXPORT_SCOPE", "El alcance de exportación no es válido.", 422)

    effective_filters = filters if scope == "filtered" else {}
    stmt = email_selection(effective_filters)
    rows = session.scalars(stmt.execution_options(yield_per=FOLLOW_UP_CHUNK_SIZE))
    started = perf_counter()
    now = datetime.now(ECUADOR_TZ)
    filename = f"SITS_Correos_Categorizados_{now:%Y%m%d_%H%M%S}.xlsx"
    temp = tempfile.NamedTemporaryFile(prefix="sits-correos-", suffix=".xlsx", delete=False)
    temp_path = temp.name
    temp.close()
    count = 0
    extended_count = 0
    follow_up_count = 0
    try:
        workbook = Workbook(write_only=True)
        correos_sheet = workbook.create_sheet("Correos")
        headings = [heading for heading, _ in CORREOS_COLUMNS]
        if include_body:
            headings.append("Cuerpo completo")
        _header(correos_sheet, headings)
        _configure_columns(correos_sheet, [22, 36, 28, 28, 42, 32, 45, 16, 16, 12, 28, 34, 22, 30, 24, 24, 28, 16, 36, 28, 28] + ([52] if include_body else []))

        extended_sheet = workbook.create_sheet("Contenido_extenso")
        _header(extended_sheet, ["ID interno del correo", "Campo", "Parte", "Total de partes", "Contenido"])
        _configure_columns(extended_sheet, [36, 26, 12, 16, 100])

        follows_sheet = None
        if include_follow_ups:
            follows_sheet = workbook.create_sheet("Seguimientos")
            _header(follows_sheet, ["ID interno del correo", "MessageId", "ID del seguimiento", "Fecha", "Detalle", "Registrado por", "Estado del requerimiento"])
            _configure_columns(follows_sheet, [36, 36, 36, 28, 100, 30, 24])

        pending: list[Correo] = []
        def write_batch(batch: list[Correo]) -> None:
            nonlocal count, extended_count, follow_up_count
            follow_ups = _follow_ups_for(session, [record.id_correo for record in batch]) if follows_sheet else {}
            for record in batch:
                values = []
                for _, field in CORREOS_COLUMNS:
                    value = getattr(record, field)
                    if isinstance(value, bool):
                        value = "Sí" if value else "No"
                    text, pieces = _append_content(value, correo_id=record.id_correo, field=field, extended_sheet=extended_sheet)
                    extended_count += pieces
                    values.append(_cell(correos_sheet, text))
                if include_body:
                    text, pieces = _append_content(record.cuerpo, correo_id=record.id_correo, field="cuerpo", extended_sheet=extended_sheet)
                    extended_count += pieces
                    values.append(_cell(correos_sheet, text))
                correos_sheet.append(values)
                count += 1
                if follows_sheet:
                    for follow in follow_ups.get(record.id_correo, []):
                        detail, pieces = _append_content(follow.detalle_seguimiento, correo_id=record.id_correo, field="detalle_seguimiento", extended_sheet=extended_sheet)
                        extended_count += pieces
                        follows_sheet.append([_cell(follows_sheet, record.id_correo), _cell(follows_sheet, record.id_externo_correo), _cell(follows_sheet, follow.id_seguimiento), _cell(follows_sheet, follow.fecha_seguimiento), _cell(follows_sheet, detail), _cell(follows_sheet, follow.seguimiento_por), _cell(follows_sheet, follow.estado_requerimiento)])
                        follow_up_count += 1

        for record in rows:
            pending.append(record)
            if len(pending) == FOLLOW_UP_CHUNK_SIZE:
                write_batch(pending)
                pending.clear()
        write_batch(pending)

        # Se fijan los rangos una vez conocidos, sin crear filas de relleno.
        correos_sheet.auto_filter.ref = f"A1:{chr(64 + len(headings))}{count + 1}" if len(headings) <= 26 else None
        extended_sheet.auto_filter.ref = f"A1:E{extended_count + 1}" if extended_count else "A1:E1"
        if follows_sheet:
            follows_sheet.auto_filter.ref = f"A1:G{follow_up_count + 1}"

        info_sheet = workbook.create_sheet("Información_exportación")
        _header(info_sheet, ["Campo", "Valor"])
        _configure_columns(info_sheet, [34, 110])
        omitted = []
        if not include_body:
            omitted.append("Cuerpo completo (no seleccionado)")
        if not include_follow_ups:
            omitted.append("Seguimientos (no seleccionados)")
        info_rows = [
            ("Fecha y hora de generación", now.isoformat()),
            ("Zona horaria", "America/Guayaquil"),
            ("Alcance", "Resultados filtrados" if scope == "filtered" else "Todos los correos autorizados"),
            ("Filtros aplicados", _active_filters(effective_filters)),
            ("Registros exportados", str(count)),
            ("Incluye cuerpo completo", "Sí" if include_body else "No"),
            ("Incluye historial de seguimientos", "Sí" if include_follow_ups else "No"),
            ("Campos omitidos", "; ".join(omitted) if omitted else "Ninguno"),
            ("Fechas", "Se exportan como texto ISO persistido, incluida su referencia de zona horaria; no se cambia el instante original."),
            ("Estados de categoría", "VALIDA, SIN_CATEGORIA, DESCONOCIDA y OTROS conservan el valor persistido. Estado de clasificación y del requerimiento se muestran en columnas distintas."),
            ("Valores vacíos", "Una celda vacía representa un valor nulo o vacío en origen; no es una clasificación inventada."),
            ("Contenido extenso", "Cuando el texto supera 32767 caracteres, la fila principal lo indica y las partes ordenadas en Contenido_extenso permiten reconstruirlo por concatenación sin separadores."),
            ("Protección contra fórmulas", "Los textos externos que empiezan con =, +, -, @, tabulador o retorno de carro se prefijan con U+200B. Si el texto original empezaba con U+200B, este se duplica. Para reconstruir, elimine un único U+200B inicial."),
            ("Correspondencia de columnas", "; ".join(f"{heading} = {field}" for heading, field in CORREOS_COLUMNS) + ("; Cuerpo completo = cuerpo" if include_body else "")),
        ]
        for key, value in info_rows:
            info_sheet.append([_cell(info_sheet, key), _cell(info_sheet, value)])
        info_sheet.auto_filter.ref = f"A1:B{len(info_rows) + 1}"

        workbook.save(temp_path)
        workbook.close()
        duration_ms = round((perf_counter() - started) * 1000)
        log_change(
            session, "correos_exportaciones", str(uuid4()), "DOWNLOAD_FILE", {},
            {"alcance": scope, "cantidad": count, "incluye_cuerpo": include_body,
             "incluye_seguimientos": include_follow_ups, "resultado": "GENERADO"},
            user.correo, "Exportación XLSX de correos", correlation_id,
        )
        return EmailExport(temp_path, filename, count, duration_ms, extended_count > 0)
    except Exception as exc:
        remove_export_file(temp_path)
        if isinstance(exc, AppError):
            raise
        raise AppError("EXPORT_GENERATION_FAILED", "No fue posible preparar el archivo de Excel.", 500) from exc
