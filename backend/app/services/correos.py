"""Ingesta segura e idempotente de correos, XLSX y seguimiento operativo."""
from datetime import date, datetime
from io import BytesIO
from time import perf_counter
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.errors import AppError
from app.core.permissions import AuthenticatedUser, authorize
from app.core.time import ECUADOR_TZ, utc_now_iso
from app.models import Correo, ErrorImportacionCorreo, LoteImportacionCorreo, SeguimientoCorreo
from app.services.audit import log_change
from app.services.correos_catalog_seed import CATEGORIAS_CORREO
from app.services.records import check_expected_version, creation_metadata, get_active, mark_updated

MAX_EMAIL_XLSX_BYTES = 50 * 1024 * 1024
MAX_PREVIEW_ROWS = 100
IMPORT_BATCH_SIZE = 500
MAX_TEXT = {"id_externo_correo": 512, "asunto": 2000, "remitente": 1000, "destinatarios": 8000, "cc": 8000, "importancia": 100, "categoria_macro": 160, "categoria_nombre": 300, "regla_disparadora": 500, "idempotency_key": 600, "cuerpo": 2_000_000}
ESTADOS_REQUERIMIENTO = {"PENDIENTE", "EN_PROCESO", "EN_ESPERA", "RESUELTO", "CERRADO"}
ESTADOS_CLASIFICACION = {"CLASIFICADO", "REVISION"}
CATEGORIAS_VALIDAS = {code for code, _ in CATEGORIAS_CORREO} - {"REVISION_MANUAL"}
ESTADO_ANALIZADO, ESTADO_CONFIRMADO = "ANALIZADO", "CONFIRMADO"

_HEADER_ALIASES = {
    "id_externo_correo": {"id", "id externo correo", "messageid", "message id"}, "asunto": {"subject", "asunto"}, "remitente": {"from", "remitente"}, "destinatarios": {"to", "destinatarios"}, "cc": {"cc"}, "fecha_recibido": {"receivedtime", "received time", "recibido en", "fecha recibido"}, "importancia": {"importance", "prioridad", "importancia"}, "cuerpo": {"body", "cuerpo"}, "tiene_adjuntos": {"hasattachments", "has attachments", "tiene adjuntos"}, "leido": {"isread", "is read", "leido"}, "categoria_macro": {"categoria macro", "category", "categoria"}, "categoria_nombre": {"categoria nombre"}, "regla_disparadora": {"regla disparadora"}, "estado_clasificacion": {"estado clasificacion"}, "idempotency_key": {"idempotency key"},
}


def max_xlsx_bytes() -> int:
    return get_settings().max_email_xlsx_bytes


def _text(value) -> str | None:
    return str(value).strip() or None if value is not None else None


def _header(value) -> str:
    return " ".join("".join(char if char.isalnum() else " " for char in str(value or "").lower()).split())


def _canonical_headers(values) -> dict[str, int]:
    result = {}
    for index, value in enumerate(values):
        for field, aliases in _HEADER_ALIASES.items():
            if _header(value) in aliases:
                result[field] = index
                break
    return result


def _value(values, positions: dict[str, int], field: str):
    index = positions.get(field)
    return values[index] if index is not None and index < len(values) else None


def _boolean(value) -> bool:
    return _header(value) in {"si", "true", "1", "yes", "x"}


def normalize_received(value) -> str | None:
    if value is None or _text(value) is None:
        return None
    if isinstance(value, datetime): result = value
    elif isinstance(value, date): result = datetime.combine(value, datetime.min.time())
    else:
        try: result = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        except ValueError as exc: raise AppError("INVALID_RECEIVED_TIME", "La fecha de recepción no es válida.", 422) from exc
    if result.tzinfo is None: result = result.replace(tzinfo=ECUADOR_TZ)
    return result.astimezone(ECUADOR_TZ).isoformat()


def _sort_key(record: dict) -> tuple[int, float | str]:
    value = record.get("fecha_recibido")
    if not value: return (0, "")
    try: return (2, datetime.fromisoformat(value).timestamp())
    except (TypeError, ValueError): return (1, str(value))


def _category(category: str | None, explicit_status) -> tuple[str, str]:
    explicit = _header(explicit_status).upper()
    if not category: return "SIN_CATEGORIA", "REVISION"
    if category == "REVISION_MANUAL": return "OTROS", "REVISION"
    if category in CATEGORIAS_VALIDAS: return "VALIDA", explicit if explicit in ESTADOS_CLASIFICACION else "CLASIFICADO"
    return "DESCONOCIDA", "REVISION"


def _error(row: int, external_id: str | None, code: str, detail: str) -> dict:
    return {"fila": row, "id_externo_correo": external_id, "codigo": code, "detalle": detail}


def _long_field(data: dict) -> str | None:
    return next((field for field, maximum in MAX_TEXT.items() if data.get(field) is not None and len(str(data[field])) > maximum), None)


def _read_xlsx(content: bytes) -> tuple[list[dict], list[dict], int, int]:
    """Procesa por filas; los errores de una fila no revocan las válidas."""
    if not content: raise AppError("EMPTY_IMPORT_FILE", "El archivo XLSX está vacío.", 422)
    if len(content) > max_xlsx_bytes(): raise AppError("IMPORT_FILE_TOO_LARGE", "El archivo XLSX supera el tamaño máximo permitido para correos.", 422)
    try: workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc: raise AppError("INVALID_XLSX", "El archivo no es un XLSX válido.", 422) from exc
    try:
        sheet = workbook["Correos_POST"] if "Correos_POST" in workbook.sheetnames else workbook.active
        header_row = None
        for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True), start=1):
            positions = _canonical_headers(row)
            if "id_externo_correo" in positions:
                header_row = number
                break
        if header_row is None: raise AppError("INVALID_IMPORT_HEADERS", "Falta el encabezado MessageId/ID requerido.", 422)
        records, errors, seen, blanks = [], [], set(), 0
        for number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(_text(value) for value in values): blanks += 1; continue
            external_id = _text(_value(values, positions, "id_externo_correo"))
            if not external_id: errors.append(_error(number, None, "MESSAGE_ID_REQUIRED", "La fila no tiene MessageId/ID.")); continue
            if external_id in seen: errors.append(_error(number, external_id, "DUPLICATE_IN_FILE", "MessageId repetido dentro del XLSX.")); continue
            seen.add(external_id)
            try:
                category = _text(_value(values, positions, "categoria_macro"))
                category_status, classification_status = _category(category, _value(values, positions, "estado_clasificacion"))
                record = {"registro_fuente": str(number), "id_externo_correo": external_id, "idempotency_key": _text(_value(values, positions, "idempotency_key")) or f"correo:{external_id}", "asunto": _text(_value(values, positions, "asunto")) or "", "remitente": _text(_value(values, positions, "remitente")), "destinatarios": _text(_value(values, positions, "destinatarios")), "cc": _text(_value(values, positions, "cc")), "fecha_recibido": normalize_received(_value(values, positions, "fecha_recibido")), "importancia": _text(_value(values, positions, "importancia")), "cuerpo": _text(_value(values, positions, "cuerpo")), "tiene_adjuntos": _boolean(_value(values, positions, "tiene_adjuntos")), "leido": _boolean(_value(values, positions, "leido")), "categoria_macro": category, "categoria_nombre": _text(_value(values, positions, "categoria_nombre")), "regla_disparadora": _text(_value(values, positions, "regla_disparadora")), "estado_clasificacion": classification_status, "estado_categoria": category_status}
            except AppError as exc: errors.append(_error(number, external_id, exc.code, str(exc.detail))); continue
            long_field = _long_field(record)
            if long_field: errors.append(_error(number, external_id, "FIELD_TOO_LONG", f"El campo {long_field} supera el tamaño permitido.")); continue
            records.append(record)
        return records, errors, len(records) + len(errors), blanks
    finally: workbook.close()


def _latest(records: list[dict]) -> list[dict]: return sorted(records, key=_sort_key, reverse=True)[:MAX_PREVIEW_ROWS]


def _snapshot(record: Correo) -> dict:
    return {"id_externo_correo": record.id_externo_correo, "asunto": record.asunto, "categoria_macro": record.categoria_macro, "estado_categoria": record.estado_categoria, "estado_clasificacion": record.estado_clasificacion, "estado_requerimiento": record.estado_requerimiento, "responsable_seguimiento": record.responsable_seguimiento}


def _create(session: Session, actor: str, data: dict, *, correlation_id: str, lote_id: str | None = None, origin: str = "MANUAL") -> Correo | None:
    existing = session.scalar(select(Correo).where(Correo.id_externo_correo == data["id_externo_correo"]))
    if existing is not None or session.scalar(select(Correo).where(Correo.idempotency_key == data["idempotency_key"])) is not None: return None
    record = Correo(id_correo=str(uuid4()), lote_id=lote_id, estado_requerimiento="PENDIENTE", archivo_fuente=origin, hoja_fuente="Correos_POST" if origin == "XLSX" else None, **data, **creation_metadata(actor))
    try:
        with session.begin_nested(): session.add(record); session.flush()
    except IntegrityError: return None
    log_change(session, "correos", record.id_correo, "CREATE", {}, _snapshot(record), actor, "Ingreso de correo", correlation_id, sensitive_record=True)
    return record


def _persist_errors(session: Session, lot_id: str, actor: str, errors: list[dict]) -> None:
    for error in errors: session.add(ErrorImportacionCorreo(id_error=str(uuid4()), lote_id=lot_id, **error, **creation_metadata(actor)))


def _chunks(values: list[dict], size: int = IMPORT_BATCH_SIZE):
    for index in range(0, len(values), size):
        yield values[index:index + size]


def _existing_import_keys(session: Session, records: list[dict]) -> tuple[set[str], set[str]]:
    """Consulta las claves existentes en grupos, sin una consulta por fila."""
    external_ids, idempotency_keys = set(), set()
    for chunk in _chunks(records):
        rows = session.execute(
            select(Correo.id_externo_correo, Correo.idempotency_key).where(
                or_(
                    Correo.id_externo_correo.in_([record["id_externo_correo"] for record in chunk]),
                    Correo.idempotency_key.in_([record["idempotency_key"] for record in chunk]),
                )
            )
        ).all()
        external_ids.update(row.id_externo_correo for row in rows)
        idempotency_keys.update(row.idempotency_key for row in rows)
    return external_ids, idempotency_keys


def _bulk_import(session: Session, actor: str, lot: LoteImportacionCorreo, records: list[dict]) -> tuple[int, int]:
    """Persiste una carga XLSX con pocos flushes y sin auditoría sensible por fila."""
    existing_external, existing_idempotency = _existing_import_keys(session, records)
    seen_external, seen_idempotency = set(existing_external), set(existing_idempotency)
    inserted, duplicates, pending = 0, 0, []
    metadata = creation_metadata(actor)
    for data in records:
        external_id, idempotency_key = data["id_externo_correo"], data["idempotency_key"]
        if external_id in seen_external or idempotency_key in seen_idempotency:
            duplicates += 1
            continue
        seen_external.add(external_id); seen_idempotency.add(idempotency_key)
        pending.append(Correo(
            id_correo=str(uuid4()), lote_id=lot.id_lote, estado_requerimiento="PENDIENTE",
            archivo_fuente="XLSX", hoja_fuente="Correos_POST", **data, **metadata,
        ))
    for chunk in _chunks(pending):
        session.add_all(chunk)
        session.flush()
    return len(pending), duplicates


def analyze_import(session: Session, user: AuthenticatedUser, *, filename: str, content: bytes, correlation_id: str):
    authorize(user, "IMPORTACION", "create"); authorize(user, "CORREOS", "create")
    if not filename.lower().endswith(".xlsx"): raise AppError("INVALID_XLSX", "El archivo debe tener extensión .xlsx.", 422)
    started = perf_counter(); records, errors, processed, blank_rows = _read_xlsx(content); classified = sum(row["estado_clasificacion"] == "CLASIFICADO" for row in records)
    lot = LoteImportacionCorreo(id_lote=str(uuid4()), nombre_archivo=filename, contenido_archivo=content, usuario_id=user.id_usuario, estado=ESTADO_ANALIZADO, total_filas=processed + blank_rows, filas_procesadas=processed, filas_clasificadas=classified, filas_revision=len(records) - classified, filas_importadas=0, filas_error=len(errors), filas_omitidas=blank_rows, filas_duplicadas=0, origen="XLSX", duracion_ms=round((perf_counter() - started) * 1000), **creation_metadata(user.correo))
    session.add(lot); session.flush(); _persist_errors(session, lot.id_lote, user.correo, errors)
    log_change(session, "lotes_importacion_correo", lot.id_lote, "IMPORT", {}, {"nombre_archivo": lot.nombre_archivo, "total_filas": lot.total_filas, "filas_procesadas": lot.filas_procesadas, "filas_error": lot.filas_error}, user.correo, "Análisis de archivo XLSX de correos", correlation_id, sensitive_record=True)
    return lot, _latest(records)


def confirm_import(session: Session, user: AuthenticatedUser, lot_id: str, *, include_revision: bool, correlation_id: str):
    authorize(user, "IMPORTACION", "create"); authorize(user, "CORREOS", "create")
    lot = get_active(session, LoteImportacionCorreo, lot_id, LoteImportacionCorreo.id_lote)
    if lot.estado == ESTADO_CONFIRMADO: raise AppError("IMPORT_ALREADY_CONFIRMED", "El lote ya fue confirmado.", 409)
    if lot.estado != ESTADO_ANALIZADO: raise AppError("INVALID_IMPORT_STATE", "El lote no está disponible para confirmación.", 409)
    started = perf_counter(); records, _, _, _ = _read_xlsx(lot.contenido_archivo); selected = records if include_revision else [row for row in records if row["estado_clasificacion"] == "CLASIFICADO"]
    if not selected: raise AppError("NO_CLASSIFIED_ROWS", "No existen correos aptos para cargar.", 422)
    inserted, duplicates = _bulk_import(session, user.correo, lot, selected)
    before = {"estado": lot.estado, "filas_importadas": lot.filas_importadas}; lot.estado, lot.filas_importadas, lot.filas_duplicadas = ESTADO_CONFIRMADO, inserted, duplicates; lot.duracion_ms = (lot.duracion_ms or 0) + round((perf_counter() - started) * 1000); mark_updated(lot, user.correo)
    log_change(session, "lotes_importacion_correo", lot.id_lote, "IMPORT", before, {"estado": lot.estado, "filas_importadas": inserted, "filas_duplicadas": duplicates, "incluyo_revision": include_revision}, user.correo, "Confirmación de carga de correos", correlation_id, sensitive_record=True)
    return lot, len(selected), inserted, duplicates


def create_integration_email(session: Session, data: dict, *, correlation_id: str, actor: str = "n8n@integration.local", origin: str = "N8N") -> tuple[Correo, bool]:
    existing = session.scalar(select(Correo).where(Correo.id_externo_correo == data["id_externo_correo"]))
    if existing is not None: return existing, False
    created = _create(session, actor, data, correlation_id=correlation_id, origin=origin)
    if created is None:
        existing = session.scalar(select(Correo).where(or_(Correo.id_externo_correo == data["id_externo_correo"], Correo.idempotency_key == data["idempotency_key"])))
        if existing is not None: return existing, False
        raise AppError("IDEMPOTENCY_CONFLICT", "No fue posible aplicar la idempotencia del correo.", 409)
    return created, True


def create_from_post(session: Session, user: AuthenticatedUser, data: dict, *, correlation_id: str) -> tuple[Correo, bool]:
    authorize(user, "CORREOS", "create")
    return create_integration_email(session, data, correlation_id=correlation_id, actor=user.correo, origin="MANUAL")


def list_emails(session: Session, user: AuthenticatedUser, *, limit: int, offset: int, filters: dict):
    authorize(user, "CORREOS", "read"); stmt = select(Correo).where(Correo.eliminado.is_(False)); values = {key: value for key, value in filters.items() if value not in (None, "")}
    if values.get("estado"): stmt = stmt.where(Correo.estado_requerimiento == values["estado"])
    for field, column in (("categoria", Correo.categoria_macro), ("origen", Correo.archivo_fuente), ("importancia", Correo.importancia)):
        if values.get(field): stmt = stmt.where(column == values[field])
    if values.get("tiene_adjuntos") is not None: stmt = stmt.where(Correo.tiene_adjuntos.is_(values["tiene_adjuntos"]))
    if values.get("fecha_desde"): stmt = stmt.where(Correo.fecha_recibido >= values["fecha_desde"])
    if values.get("fecha_hasta"): stmt = stmt.where(Correo.fecha_recibido <= values["fecha_hasta"] + "T23:59:59")
    for field, column in (("asunto", Correo.asunto), ("remitente", Correo.remitente), ("destinatario", Correo.destinatarios), ("message_id", Correo.id_externo_correo)):
        if values.get(field): stmt = stmt.where(column.ilike(f"%{values[field]}%"))
    if values.get("texto"):
        term = f"%{values['texto']}%"; stmt = stmt.where(or_(Correo.asunto.ilike(term), Correo.remitente.ilike(term), Correo.destinatarios.ilike(term), Correo.id_externo_correo.ilike(term)))
    total = session.scalar(select(func.count()).select_from(stmt.subquery())) or 0; order = Correo.asunto.asc() if values.get("orden") == "asunto_asc" else Correo.fecha_recibido.desc()
    return session.scalars(stmt.order_by(order, Correo.fecha_creacion.desc()).offset(offset).limit(limit)).all(), total


def summary(session: Session, user: AuthenticatedUser) -> dict:
    authorize(user, "CORREOS", "read"); base = Correo.eliminado.is_(False)
    def grouped(column): return [{"valor": value or "Sin información", "total": total} for value, total in session.execute(select(column, func.count()).where(base).group_by(column).order_by(func.count().desc()).limit(10))]
    return {"total": session.scalar(select(func.count()).select_from(Correo).where(base)) or 0, "pendientes": session.scalar(select(func.count()).select_from(Correo).where(base, Correo.estado_requerimiento == "PENDIENTE")) or 0, "en_seguimiento": session.scalar(select(func.count()).select_from(Correo).where(base, Correo.estado_requerimiento == "EN_PROCESO")) or 0, "sin_clasificar": session.scalar(select(func.count()).select_from(Correo).where(base, Correo.estado_categoria != "VALIDA")) or 0, "por_categoria": grouped(Correo.categoria_macro), "por_origen": grouped(Correo.archivo_fuente), "principales_remitentes": grouped(Correo.remitente)}


def list_lots(session: Session, user: AuthenticatedUser) -> list[LoteImportacionCorreo]:
    authorize(user, "IMPORTACION", "read"); return list(session.scalars(select(LoteImportacionCorreo).where(LoteImportacionCorreo.eliminado.is_(False)).order_by(LoteImportacionCorreo.fecha_creacion.desc()).limit(50)))


def list_lot_errors(session: Session, user: AuthenticatedUser, lot_id: str) -> list[ErrorImportacionCorreo]:
    authorize(user, "IMPORTACION", "read"); get_active(session, LoteImportacionCorreo, lot_id, LoteImportacionCorreo.id_lote)
    return list(session.scalars(select(ErrorImportacionCorreo).where(ErrorImportacionCorreo.lote_id == lot_id, ErrorImportacionCorreo.eliminado.is_(False)).order_by(ErrorImportacionCorreo.fila).limit(500)))


def get_detail(session: Session, user: AuthenticatedUser, correo_id: str, *, correlation_id: str) -> tuple[Correo, list[SeguimientoCorreo]]:
    authorize(user, "CORREOS", "read", sensitive=True); record = get_active(session, Correo, correo_id, Correo.id_correo); log_change(session, "correos", record.id_correo, "VIEW_SENSITIVE", {}, {"estado_requerimiento": record.estado_requerimiento}, user.correo, "Consulta del cuerpo de correo", correlation_id, sensitive_record=True)
    follows = session.scalars(select(SeguimientoCorreo).where(SeguimientoCorreo.correo_id == record.id_correo, SeguimientoCorreo.eliminado.is_(False)).order_by(SeguimientoCorreo.fecha_seguimiento.desc(), SeguimientoCorreo.fecha_creacion.desc())).all()
    return record, list(follows)


def add_follow_up(session: Session, user: AuthenticatedUser, correo_id: str, *, expected_version: int, detail: str, responsible: str | None, state: str, correlation_id: str) -> tuple[Correo, SeguimientoCorreo]:
    authorize(user, "CORREOS", "edit", sensitive=True)
    if state not in ESTADOS_REQUERIMIENTO or not _text(detail): raise AppError("INVALID_INPUT", "El detalle y el estado del seguimiento son obligatorios y válidos.", 422)
    record = get_active(session, Correo, correo_id, Correo.id_correo); check_expected_version(record, expected_version); before = _snapshot(record)
    follow = SeguimientoCorreo(id_seguimiento=str(uuid4()), correo_id=record.id_correo, fecha_seguimiento=utc_now_iso(), detalle_seguimiento=detail.strip(), seguimiento_por=_text(responsible) or user.nombre, estado_requerimiento=state, **creation_metadata(user.correo)); session.add(follow); record.estado_requerimiento, record.responsable_seguimiento = state, follow.seguimiento_por; mark_updated(record, user.correo); session.flush(); log_change(session, "correos", record.id_correo, "UPDATE", before, _snapshot(record), user.correo, "Registro de seguimiento", correlation_id, sensitive_record=True)
    return record, follow
