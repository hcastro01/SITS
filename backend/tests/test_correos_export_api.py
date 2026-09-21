"""Contrato HTTP de la exportación XLSX de Correos.

Estas pruebas protegen la ruta estática ``/exportar`` frente a la ruta dinámica
``/{correo_id}``: un POST autorizado debe generar el archivo, nunca terminar en
un 405 por ser tratado como el detalle de un correo.
"""

import unittest
from io import BytesIO
from tempfile import TemporaryDirectory

from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.api.deps import get_db
import app.db.session as db_session
from app.db.session import build_engine
from app.main import app
from app.models import User
from app.core.permissions import resolve_current_user
from app.services.correos import add_follow_up, create_from_post
from app.services.security_seed import seed_security


class CorreosExportApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.engine = build_engine(f"sqlite:///{self.directory.name}/correos-export.db")
        self.original_engine = db_session.engine
        db_session.engine = self.engine
        command.upgrade(Config("alembic.ini"), "head")
        with Session(self.engine) as session, session.begin():
            seed_security(session)
            session.add_all([
                User(id_usuario="admin", correo="admin@example.com", nombre="Admin", rol_id="ROLE_ADMIN", estado="ACTIVO"),
                User(id_usuario="consulta", correo="consulta@example.com", nombre="Consulta", rol_id="ROLE_CONSULTA", estado="ACTIVO"),
            ])
            session.flush()
            admin = resolve_current_user(session, "admin@example.com")
            first, _ = create_from_post(session, admin, self._correo("mail-1", "ana@example.test", "=Fórmula externa"), correlation_id="setup")
            second, _ = create_from_post(session, admin, self._correo("mail-2", "luis@example.test", "Contenido normal"), correlation_id="setup")
            add_follow_up(session, admin, first.id_correo, expected_version=first.version,
                          detail="Seguimiento de prueba", responsible="Admin", state="EN_PROCESO",
                          correlation_id="setup")
            self.assertNotEqual(first.id_correo, second.id_correo)

        def override_db():
            with Session(self.engine) as session:
                try:
                    yield session
                    session.commit()
                except Exception:
                    session.rollback()
                    raise

        app.dependency_overrides[get_db] = override_db
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        db_session.engine = self.original_engine
        self.engine.dispose()
        self.directory.cleanup()

    @staticmethod
    def _correo(message_id: str, sender: str, body: str) -> dict:
        return {
            "id_externo_correo": message_id,
            "idempotency_key": f"correo:{message_id}",
            "asunto": f"Asunto {message_id}",
            "remitente": sender,
            "destinatarios": "social@example.test",
            "cc": None,
            "fecha_recibido": "2026-09-21T10:00:00-05:00",
            "importancia": "normal",
            "cuerpo": body,
            "tiene_adjuntos": False,
            "leido": False,
            "categoria_macro": "CASOS_TALENTO_HUMANO",
            "categoria_nombre": "Casos de talento humano",
            "regla_disparadora": "caso",
            "estado_clasificacion": "CLASIFICADO",
        }

    def _login_admin(self):
        self.assertEqual(self.client.post("/api/v1/auth/login", json={"correo": "admin@example.com"}).status_code, 200)

    def _login_consulta(self):
        self.assertEqual(self.client.post("/api/v1/auth/login", json={"correo": "consulta@example.com"}).status_code, 200)

    @staticmethod
    def _workbook(response):
        return load_workbook(BytesIO(response.content), read_only=True, data_only=False, keep_links=False)

    def test_post_exportar_returns_a_valid_full_xlsx_without_sensitive_content(self):
        self._login_admin()
        response = self.client.post("/api/v1/correos/exportar", json={
            "alcance": "all",
            "filtros": {"remitente": "does-not-limit-all@example.test"},
            "incluir_cuerpo": False,
            "incluir_seguimientos": False,
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment;", response.headers["content-disposition"])
        self.assertIn(".xlsx", response.headers["content-disposition"])
        workbook = self._workbook(response)
        try:
            self.assertEqual(workbook.sheetnames, ["Correos", "Contenido_extenso", "Información_exportación"])
            rows = list(workbook["Correos"].values)
            self.assertEqual([row[1] for row in rows[1:]], ["mail-2", "mail-1"])
            self.assertNotIn("Cuerpo completo", rows[0])
            self.assertIn("Categoría", rows[0])
            self.assertIn("Estado de clasificación", rows[0])
        finally:
            workbook.close()

    def test_filtered_export_includes_only_requested_sensitive_content_and_follow_ups(self):
        self._login_admin()
        response = self.client.post("/api/v1/correos/exportar", json={
            "alcance": "filtered",
            "filtros": {"remitente": "ana@example.test", "orden": "recibido_desc"},
            "incluir_cuerpo": True,
            "incluir_seguimientos": True,
        })

        self.assertEqual(response.status_code, 200)
        workbook = self._workbook(response)
        try:
            self.assertEqual(workbook.sheetnames, ["Correos", "Contenido_extenso", "Seguimientos", "Información_exportación"])
            rows = list(workbook["Correos"].values)
            headers = rows[0]
            self.assertEqual([row[1] for row in rows[1:]], ["mail-1"])
            self.assertTrue(rows[1][headers.index("Cuerpo completo")].startswith("\u200b="))
            self.assertEqual(len(list(workbook["Seguimientos"].values)) - 1, 1)
        finally:
            workbook.close()

    def test_export_requires_session_and_returns_useful_domain_errors(self):
        unauthenticated = TestClient(app).post("/api/v1/correos/exportar", json={"alcance": "all", "filtros": {}})
        self.assertEqual(unauthenticated.status_code, 401)
        self._login_admin()
        empty = self.client.post("/api/v1/correos/exportar", json={
            "alcance": "filtered", "filtros": {"remitente": "nadie@example.test"},
        })
        self.assertEqual((empty.status_code, empty.json()["code"]), (422, "NO_EXPORT_RESULTS"))

    def test_export_requires_the_correos_permission(self):
        self._login_consulta()
        response = self.client.post("/api/v1/correos/exportar", json={"alcance": "all", "filtros": {}})
        self.assertEqual((response.status_code, response.json()["code"]), (403, "FORBIDDEN"))

    def test_full_export_is_not_limited_by_the_dashboard_page_size(self):
        with Session(self.engine) as session, session.begin():
            admin = resolve_current_user(session, "admin@example.com")
            for index in range(101):
                create_from_post(
                    session,
                    admin,
                    self._correo(f"extra-{index:03d}", "lote@example.test", "Carga de prueba"),
                    correlation_id="large-export",
                )

        self._login_admin()
        response = self.client.post("/api/v1/correos/exportar", json={"alcance": "all", "filtros": {}})

        self.assertEqual(response.status_code, 200)
        workbook = self._workbook(response)
        try:
            self.assertEqual(len(list(workbook["Correos"].values)) - 1, 103)
        finally:
            workbook.close()

    def test_export_preflight_allows_the_post_contract(self):
        response = self.client.options("/api/v1/correos/exportar", headers={
            "Origin": "http://localhost:8081",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("POST", response.headers["access-control-allow-methods"])


if __name__ == "__main__":
    unittest.main()
