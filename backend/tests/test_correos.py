import unittest
from io import BytesIO
from tempfile import TemporaryDirectory

from alembic import command
from alembic.config import Config
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

import app.db.session as db_session
from app.core.errors import AppError
from app.core.permissions import resolve_current_user
from app.db.session import build_engine
from app.models import Auditoria, Correo, ErrorImportacionCorreo, Permission, SeguimientoCorreo, User
from app.services.correos import add_follow_up, analyze_import, confirm_import, create_from_post
from app.services.correos_export import generate_email_export, remove_export_file
from app.services.security_seed import seed_security


HEADERS = ["ID", "Subject", "From", "ReceivedTime", "Body", "categoria_macro", "categoria_nombre", "estado_clasificacion", "n8n_enviar_post", "n8n_estado_envio"]


def xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Correos_POST"
    sheet.append(["Carga de prueba"])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class CorreosServiceTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.engine = build_engine(f"sqlite:///{self.directory.name}/test.db")
        self.original_engine = db_session.engine
        db_session.engine = self.engine
        command.upgrade(Config("alembic.ini"), "head")
        with Session(self.engine) as session, session.begin():
            seed_security(session)
            session.add_all([
                User(id_usuario="admin", correo="admin@example.com", nombre="Admin", rol_id="ROLE_ADMIN", estado="ACTIVO"),
                User(id_usuario="consulta", correo="consulta@example.com", nombre="Consulta", rol_id="ROLE_CONSULTA", estado="ACTIVO"),
            ])

    def tearDown(self):
        db_session.engine = self.original_engine
        self.engine.dispose()
        self.directory.cleanup()

    def user(self, session, email="admin@example.com"):
        return resolve_current_user(session, email)

    def content(self):
        return xlsx([
            ["mail-1", "Permiso", "ana@example.com", "2026-09-20T10:00:00", "Contenido privado 1", "PERMISOS_VACACIONES_LICENCIAS", "Permisos", "CLASIFICADO", "SI", "PENDIENTE"],
            ["mail-2", "Sin patrón", "luis@example.com", "2026-09-21T10:00:00", "Contenido privado 2", "REVISION_MANUAL", "Revisión manual", "REVISION", "SI", "PENDIENTE"],
            ["mail-3", "Omitido", "ana@example.com", "2026-09-22T10:00:00", "No se carga", "SALUD_OCUPACIONAL", "Salud", "CLASIFICADO", "NO", "PENDIENTE"],
        ])

    def test_analysis_detects_unclassified_and_preview_orders_latest(self):
        with Session(self.engine) as session, session.begin():
            lot, preview = analyze_import(session, self.user(session), filename="correos.xlsx", content=self.content(), correlation_id="test")
            self.assertEqual((lot.total_filas, lot.filas_clasificadas, lot.filas_revision), (3, 2, 1))
            self.assertEqual(preview[0]["id_externo_correo"], "mail-3")
            self.assertEqual(session.scalars(select(Correo)).all(), [])

    def test_confirmation_excludes_revision_when_user_declines_it(self):
        with Session(self.engine) as session, session.begin():
            lot, _ = analyze_import(session, self.user(session), filename="correos.xlsx", content=self.content(), correlation_id="test")
            confirmed, selected, inserted, duplicates = confirm_import(session, self.user(session), lot.id_lote, include_revision=False, correlation_id="test")
            rows = session.scalars(select(Correo)).all()
            self.assertEqual((confirmed.estado, selected, inserted, duplicates, len(rows)), ("CONFIRMADO", 2, 2, 0, 2))
            self.assertEqual(rows[0].id_externo_correo, "mail-1")

    def test_post_is_idempotent_and_follow_up_updates_state(self):
        payload = {
            "id_externo_correo": "mail-api", "idempotency_key": "correo:mail-api", "asunto": "Caso",
            "remitente": "ana@example.com", "destinatarios": None, "cc": None, "fecha_recibido": "2026-09-20T10:00:00",
            "importancia": None, "cuerpo": "Contenido privado", "tiene_adjuntos": False, "leido": False,
            "categoria_macro": "CASOS_TALENTO_HUMANO", "categoria_nombre": "Casos de talento humano",
            "regla_disparadora": "caso", "estado_clasificacion": "CLASIFICADO",
        }
        with Session(self.engine) as session, session.begin():
            record, created = create_from_post(session, self.user(session), payload, correlation_id="test")
            existing, created_again = create_from_post(session, self.user(session), payload, correlation_id="test")
            updated, follow = add_follow_up(session, self.user(session), record.id_correo, expected_version=record.version, detail="Se contactó al remitente", responsible="Admin", state="EN_PROCESO", correlation_id="test")
            self.assertTrue(created)
            self.assertFalse(created_again)
            self.assertEqual(existing.id_correo, record.id_correo)
            self.assertEqual((updated.estado_requerimiento, follow.estado_requerimiento), ("EN_PROCESO", "EN_PROCESO"))
            self.assertEqual(len(session.scalars(select(SeguimientoCorreo)).all()), 1)

    def test_consultation_role_cannot_analyze_file(self):
        with Session(self.engine) as session, session.begin(), self.assertRaises(AppError) as context:
            analyze_import(session, self.user(session, "consulta@example.com"), filename="correos.xlsx", content=self.content(), correlation_id="test")
        self.assertEqual(context.exception.code, "FORBIDDEN")

    def test_invalid_rows_are_traced_without_discarding_valid_rows(self):
        content = xlsx([
            ["ok-1", "Asunto", "ana@example.com", "2026-09-20T10:00:00", "Cuerpo", "CASOS_TALENTO_HUMANO", "Casos", "CLASIFICADO", "SI", "PENDIENTE"],
            ["", "Sin identificador", "ana@example.com", "2026-09-20T10:00:00", "Cuerpo", "CASOS_TALENTO_HUMANO", "Casos", "CLASIFICADO", "SI", "PENDIENTE"],
            ["bad-date", "Fecha", "ana@example.com", "no es fecha", "Cuerpo", "CASOS_TALENTO_HUMANO", "Casos", "CLASIFICADO", "SI", "PENDIENTE"],
            ["ok-1", "Duplicado", "ana@example.com", "2026-09-20T10:00:00", "Cuerpo", "CASOS_TALENTO_HUMANO", "Casos", "CLASIFICADO", "SI", "PENDIENTE"],
        ])
        with Session(self.engine) as session, session.begin():
            lot, _ = analyze_import(session, self.user(session), filename="correos.xlsx", content=content, correlation_id="test")
            errors = session.scalars(select(ErrorImportacionCorreo).where(ErrorImportacionCorreo.lote_id == lot.id_lote)).all()
            self.assertEqual((lot.total_filas, lot.filas_procesadas, lot.filas_error, len(errors)), (4, 4, 3, 3))
            self.assertEqual({error.codigo for error in errors}, {"MESSAGE_ID_REQUIRED", "INVALID_RECEIVED_TIME", "DUPLICATE_IN_FILE"})
            _, selected, inserted, duplicates = confirm_import(session, self.user(session), lot.id_lote, include_revision=True, correlation_id="test")
            self.assertEqual((selected, inserted, duplicates), (1, 1, 0))

    def test_database_deduplication_is_visible_in_import_result(self):
        with Session(self.engine) as session, session.begin():
            first, _ = analyze_import(session, self.user(session), filename="first.xlsx", content=self.content(), correlation_id="first")
            confirm_import(session, self.user(session), first.id_lote, include_revision=True, correlation_id="first")
            second, _ = analyze_import(session, self.user(session), filename="second.xlsx", content=self.content(), correlation_id="second")
            _, selected, inserted, duplicates = confirm_import(session, self.user(session), second.id_lote, include_revision=True, correlation_id="second")
            self.assertEqual((selected, inserted, duplicates), (3, 0, 3))
            self.assertEqual(len(session.scalars(select(Correo)).all()), 3)

    def test_large_confirmation_uses_batched_persistence(self):
        rows = [
            [f"mail-{index}", "Caso", "ana@example.com", "2026-09-20T10:00:00", "Contenido", "CASOS_TALENTO_HUMANO", "Casos", "CLASIFICADO", "SI", "PENDIENTE"]
            for index in range(1_001)
        ]
        with Session(self.engine) as session, session.begin():
            lot, _ = analyze_import(session, self.user(session), filename="large.xlsx", content=xlsx(rows), correlation_id="test")
            confirmed, selected, inserted, duplicates = confirm_import(session, self.user(session), lot.id_lote, include_revision=True, correlation_id="test")
            self.assertEqual((confirmed.estado, selected, inserted, duplicates), ("CONFIRMADO", 1_001, 1_001, 0))
            self.assertEqual(len(session.scalars(select(Correo)).all()), 1_001)

    def test_xlsx_export_preserves_selection_content_and_permissions(self):
        with Session(self.engine) as session, session.begin():
            lot, _ = analyze_import(session, self.user(session), filename="correos.xlsx", content=self.content(), correlation_id="export-setup")
            confirm_import(session, self.user(session), lot.id_lote, include_revision=True, correlation_id="export-setup")
            first = session.scalar(select(Correo).where(Correo.id_externo_correo == "mail-1"))
            assert first is not None
            first.cuerpo = "=SUM(A1:A2)\n" + ("ñ" * 33_000)
            first.destinatarios = "ana@example.test; José <jose@example.test>"
            add_follow_up(session, self.user(session), first.id_correo, expected_version=first.version,
                          detail="@seguimiento\n" + ("é" * 33_000), responsible="Admin",
                          state="EN_PROCESO", correlation_id="export-follow")
            result = generate_email_export(
                session, self.user(session), scope="filtered", filters={"remitente": "ana@example.com", "orden": "recibido_desc"},
                include_body=True, include_follow_ups=True, correlation_id="export-test",
            )
            self.assertEqual(result.count, 2)
            self.assertTrue(result.extended_content)
            self.assertTrue(result.filename.startswith("SITS_Correos_Categorizados_"))
            workbook = load_workbook(result.path, read_only=True, data_only=False, keep_links=False)
            try:
                self.assertEqual(workbook.sheetnames, ["Correos", "Contenido_extenso", "Seguimientos", "Información_exportación"])
                correos = list(workbook["Correos"].values)
                headers = correos[0]
                self.assertIn("Estado de categoría", headers)
                self.assertIn("Estado de clasificación", headers)
                self.assertIn("Estado del requerimiento", headers)
                self.assertIn("Cuerpo completo", headers)
                self.assertEqual([row[1] for row in correos[1:]], ["mail-3", "mail-1"])
                body_index = headers.index("Cuerpo completo")
                self.assertEqual(correos[2][body_index], "[Contenido extendido: consulte Contenido_extenso]")
                parts = [row for row in list(workbook["Contenido_extenso"].values)[1:] if row[1] == "cuerpo"]
                rebuilt = "".join(str(row[4]).removeprefix("\u200b") for row in sorted(parts, key=lambda row: row[2]))
                self.assertEqual(rebuilt, first.cuerpo)
                self.assertEqual(len(list(workbook["Seguimientos"].values)) - 1, 1)
                info = dict(list(workbook["Información_exportación"].values)[1:])
                self.assertEqual(info["Registros exportados"], "2")
                self.assertIn("U+200B", info["Protección contra fórmulas"])
            finally:
                workbook.close()
            self.assertTrue(session.scalars(select(Auditoria).where(Auditoria.tabla == "correos_exportaciones", Auditoria.accion == "DOWNLOAD_FILE")).first())
            remove_export_file(result.path)

            permission = session.scalar(select(Permission).where(Permission.rol_id == "ROLE_CONSULTA", Permission.modulo == "CORREOS"))
            assert permission is not None
            permission.puede_exportar = True
            session.flush()
            self.assertTrue(self.user(session, "consulta@example.com").permisos["CORREOS"]["export"])
            with self.assertRaises(AppError) as context:
                generate_email_export(session, self.user(session, "consulta@example.com"), scope="all", filters={},
                                      include_body=True, include_follow_ups=False, correlation_id="export-forbidden")
            self.assertEqual(context.exception.code, "FORBIDDEN")


if __name__ == "__main__":
    unittest.main()
